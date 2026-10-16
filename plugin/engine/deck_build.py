#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把课件规格（JSON）渲染成 pptx。

设计取向是"讲义式"而不是"发布会式"：
  * 原文是主角。quote_side 版式把演示段落放在整幅左侧，字号最大、行距最松；
    作者的说明与批注退到右侧窄栏，字号明显小一档——学生的视线先落在原文上。
  * 不做装饰性图形。没有图标、没有色块拼贴，只有一条分隔线和一个页码。
  * 标题写具体的东西（"兔头、香椿与一碗野蕈粉"），不写"从X到Y""全流程"这类
    总结腔——课件标题的作用是告诉学生这一页在讲哪一段文字。

版式：
  cover      封面
  section    章隔页
  bullets    条列
  quote_side 演示原文（主体）+ 说明／批注（旁栏）  ← 主要版式
  two_quote  两段原文并置对照
  table      表格
  end        结束页

JSON 里每个 slide 形如：
  {"layout": "quote_side", "title": "...", "kicker": "《安乐居》·汪曾祺",
   "quote": "原文……", "source": "《大师们的写作课》p28",
   "side": [{"label": "作者的说明", "text": "..."}, {"label": "批注", "text": "..."}]}

用法：
  python deck_build.py --spec deck.json --out deck.pptx
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------- 视觉常量
BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK_SOFT = RGBColor(0x4A, 0x4A, 0x4A)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
ACCENT = RGBColor(0x1F, 0x4E, 0x63)
ACCENT_SOFT = RGBColor(0x6E, 0x8B, 0x99)
RULE = RGBColor(0xD6, 0xDA, 0xDC)
PAPER = RGBColor(0xF7, 0xF6, 0xF3)

FONT_SERIF = "宋体"        # 原文
FONT_SANS = "微软雅黑"      # 标题、旁栏、页脚

W, H = 13.333, 7.5
MARGIN = 0.72


def measure(text: str, size_pt: float) -> float:
    """按字号估算字符串宽度（英寸）。中日韩字符按全宽算，西文按 0.55 宽算。"""
    units = 0.0
    for ch in text:
        units += 1.0 if ord(ch) > 0x2E80 else 0.55
    return units * size_pt / 72.0


# 估算与实际渲染总有出入（字体度量、行高取整），留出余量比事后压缩安全。
# 这个系数是拿 deck_preview.py 的实测结果反推出来的：1.0 时 29 页里有 6 处溢出。
SAFETY = 1.22


def fit(text: str, box_w: float, box_h: float, max_pt: float, min_pt: float, line_spacing: float = 1.4) -> float:
    """从 max_pt 往下试，找到能塞进 box 的字号。"""
    size = max_pt
    while size > min_pt:
        lines = 0
        for para in text.split("\n"):
            width = max(0.1, box_w)
            lines += max(1, math.ceil(measure(para, size) / width))
        needed = lines * size * line_spacing / 72.0 * SAFETY
        if needed <= box_h:
            return size
        size -= 0.5
    return min_pt


def add_box(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return frame


def write(frame, runs, *, size, color=INK, font=FONT_SANS, bold=False, spacing=1.4,
          align=PP_ALIGN.LEFT, space_after=0, first=False):
    para = frame.paragraphs[0] if first else frame.add_paragraph()
    para.alignment = align
    para.line_spacing = spacing
    if space_after:
        para.space_after = Pt(space_after)
    run = para.add_run()
    run.text = runs
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    return para


def rule(slide, x, y, w, color=RULE, thickness=0.9):
    line = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Pt(thickness))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return slide


def header(slide, kicker: str | None, title: str, *, title_size=23.0):
    y = 0.52
    if kicker:
        frame = add_box(slide, MARGIN, y, W - 2 * MARGIN, 0.3)
        write(frame, kicker, size=11.5, color=ACCENT_SOFT, font=FONT_SANS, spacing=1.0, first=True)
        y += 0.32
    size = fit(title, W - 2 * MARGIN, 0.62, title_size, 15.0, 1.15)
    frame = add_box(slide, MARGIN, y, W - 2 * MARGIN, 0.62)
    write(frame, title, size=size, color=ACCENT, font=FONT_SANS, bold=True, spacing=1.15, first=True)
    return y + 0.78


def footer(slide, page: int, note: str | None = None):
    rule(slide, MARGIN, H - 0.62, W - 2 * MARGIN)
    frame = add_box(slide, MARGIN, H - 0.5, W - 2 * MARGIN - 0.7, 0.3)
    write(frame, note or "", size=8.5, color=MUTED, spacing=1.0, first=True)
    frame = add_box(slide, W - MARGIN - 0.7, H - 0.5, 0.7, 0.3)
    write(frame, str(page), size=8.5, color=MUTED, spacing=1.0, align=PP_ALIGN.RIGHT, first=True)


# ----------------------------------------------------------------- 版式
def slide_cover(prs, spec, page):
    slide = new_slide(prs)
    y = 2.15
    if spec.get("kicker"):
        frame = add_box(slide, MARGIN + 0.05, y - 0.5, W - 2 * MARGIN, 0.35)
        write(frame, spec["kicker"], size=12, color=ACCENT_SOFT, spacing=1.0, first=True)
    frame = add_box(slide, MARGIN, y, W - 2 * MARGIN - 1.2, 1.5)
    size = fit(spec["title"], W - 2 * MARGIN - 1.2, 1.5, 40.0, 22.0, 1.12)
    write(frame, spec["title"], size=size, color=ACCENT, bold=True, spacing=1.12, first=True)
    rule(slide, MARGIN, y + 1.62, 1.5, ACCENT, 1.4)
    if spec.get("subtitle"):
        frame = add_box(slide, MARGIN, y + 1.95, W - 2 * MARGIN - 1.6, 1.1)
        write(frame, spec["subtitle"], size=15, color=INK_SOFT, font=FONT_SERIF, spacing=1.5, first=True)
    if spec.get("lines"):
        frame = add_box(slide, MARGIN, H - 1.9, W - 2 * MARGIN, 1.1)
        for i, line in enumerate(spec["lines"]):
            write(frame, line, size=11, color=MUTED, spacing=1.45, first=(i == 0))
    return slide


def slide_section(prs, spec, page):
    slide = new_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    frame = add_box(slide, MARGIN, 2.5, W - 2 * MARGIN - 1.5, 1.3)
    size = fit(spec["title"], W - 2 * MARGIN - 1.5, 1.3, 32.0, 20.0, 1.15)
    write(frame, spec["title"], size=size, color=ACCENT, bold=True, spacing=1.15, first=True)
    rule(slide, MARGIN, 3.95, 1.4, ACCENT, 1.4)
    if spec.get("subtitle"):
        frame = add_box(slide, MARGIN, 4.25, W - 2 * MARGIN - 2.0, 1.2)
        write(frame, spec["subtitle"], size=14, color=INK_SOFT, font=FONT_SERIF, spacing=1.5, first=True)
    return slide


def slide_bullets(prs, spec, page):
    slide = new_slide(prs)
    y = header(slide, spec.get("kicker"), spec["title"])
    body_h = H - y - 0.95
    bullets = spec.get("bullets") or []
    text = "\n".join(bullets)
    size = fit(text, W - 2 * MARGIN - 0.35, body_h, 15.5, 10.5, 1.65)
    frame = add_box(slide, MARGIN, y, W - 2 * MARGIN, body_h)
    for i, bullet in enumerate(bullets):
        write(frame, bullet, size=size, color=INK_SOFT, spacing=1.65, space_after=size * 0.55, first=(i == 0))
    footer(slide, page, spec.get("source"))
    return slide


def est_height(text: str, box_w: float, size: float, spacing: float = 1.56) -> float:
    """旁栏文本框的高度预算。spacing 比实际渲染用的 1.45 略大，
    是因为字号换算成像素时会向上取整，行高会再涨一点。"""
    lines = 0
    for para in text.split("\n"):
        lines += max(1, math.ceil(measure(para, size) / max(0.1, box_w)))
    return lines * size * spacing / 72.0 * SAFETY


def slide_quote_side(prs, spec, page):
    """演示原文占主体，说明与批注在右侧窄栏。整套课件的核心版式。"""
    slide = new_slide(prs)
    y = header(slide, spec.get("kicker"), spec["title"])

    main_w = 7.75
    side_x = MARGIN + main_w + 0.5
    side_w = W - MARGIN - side_x
    body_h = H - y - 0.95

    quote = spec.get("quote", "")
    frame = add_box(slide, MARGIN, y, main_w, body_h - 0.32)
    size = fit(quote, main_w, body_h - 0.32, 15.0, 9.0, 1.5)
    write(frame, quote, size=size, color=INK, font=FONT_SERIF, spacing=1.5, first=True)
    if spec.get("source"):
        frame = add_box(slide, MARGIN, y + body_h - 0.26, main_w, 0.26)
        write(frame, spec["source"], size=9.5, color=MUTED, spacing=1.0, first=True)

    rule(slide, side_x - 0.28, y, 0.9, RULE, 0.9)

    # 旁栏整体先试排一遍：块多、字长就统一下调字号，避免溢出页面
    blocks = spec.get("side") or []
    side_size = 10.0
    for candidate in (10.0, 9.5, 9.0, 8.5, 8.0, 7.5):
        total = sum(0.3 + est_height(b["text"], side_w, candidate) + 0.3 for b in blocks)
        side_size = candidate
        if total <= body_h:
            break

    cursor = y
    for block in blocks:
        frame = add_box(slide, side_x, cursor, side_w, 0.28)
        write(frame, block["label"], size=side_size + 0.5, color=ACCENT, bold=True, spacing=1.0, first=True)
        cursor += 0.3
        height = est_height(block["text"], side_w, side_size)
        frame = add_box(slide, side_x, cursor, side_w, height + 0.1)
        write(frame, block["text"], size=side_size, color=INK_SOFT, spacing=1.45, first=True)
        cursor += height + 0.3

    footer(slide, page, spec.get("footer"))
    return slide


def slide_two_quote(prs, spec, page):
    slide = new_slide(prs)
    y = header(slide, spec.get("kicker"), spec["title"])
    cols = spec.get("columns") or []
    gap = 0.5
    col_w = (W - 2 * MARGIN - gap * (len(cols) - 1)) / max(1, len(cols))
    body_h = H - y - 0.95
    for i, col in enumerate(cols):
        x = MARGIN + i * (col_w + gap)
        frame = add_box(slide, x, y, col_w, 0.34)
        write(frame, col.get("heading", ""), size=11, color=ACCENT, bold=True, spacing=1.05, first=True)
        top = y + 0.42
        if col.get("source"):
            frame = add_box(slide, x, top, col_w, 0.24)
            write(frame, col["source"], size=9, color=MUTED, spacing=1.0, first=True)
            top += 0.28
        text = col.get("quote", "")
        quote_h = y + body_h - 0.42 - top
        size = fit(text, col_w, quote_h, 13.0, 8.5, 1.45)
        frame = add_box(slide, x, top, col_w, quote_h)
        write(frame, text, size=size, color=INK, font=FONT_SERIF, spacing=1.45, first=True)
        if col.get("note"):
            frame = add_box(slide, x, y + body_h - 0.42, col_w, 0.42)
            write(frame, col["note"], size=9.5, color=MUTED, spacing=1.25, first=True)
    footer(slide, page, spec.get("footer"))
    return slide


def slide_table(prs, spec, page):
    slide = new_slide(prs)
    y = header(slide, spec.get("kicker"), spec["title"])
    rows = spec["rows"]
    cols = len(rows[0])
    width = W - 2 * MARGIN
    height = min(H - y - 1.0, 0.55 + 0.62 * (len(rows) - 1))
    shape = slide.shapes.add_table(len(rows), cols, Inches(MARGIN), Inches(y), Inches(width), Inches(height))
    table = shape.table
    table.first_row = True
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG if r else PAPER
            frame = cell.text_frame
            frame.word_wrap = True
            para = frame.paragraphs[0]
            para.line_spacing = 1.25
            run = para.add_run()
            run.text = str(value)
            run.font.size = Pt(11.5 if r == 0 else 10.5)
            run.font.name = FONT_SANS
            run.font.bold = r == 0
            run.font.color.rgb = ACCENT if r == 0 else INK_SOFT
    footer(slide, page, spec.get("footer"))
    return slide


def slide_end(prs, spec, page):
    slide = new_slide(prs)
    frame = add_box(slide, MARGIN, 2.7, W - 2 * MARGIN - 1.5, 1.2)
    size = fit(spec["title"], W - 2 * MARGIN - 1.5, 1.2, 30.0, 18.0, 1.15)
    write(frame, spec["title"], size=size, color=ACCENT, bold=True, spacing=1.15, first=True)
    rule(slide, MARGIN, 4.05, 1.4, ACCENT, 1.4)
    if spec.get("lines"):
        frame = add_box(slide, MARGIN, 4.35, W - 2 * MARGIN - 1.5, 2.0)
        for i, line in enumerate(spec["lines"]):
            write(frame, line, size=11.5, color=INK_SOFT, spacing=1.55, first=(i == 0))
    return slide


RENDERERS = {
    "cover": slide_cover,
    "section": slide_section,
    "bullets": slide_bullets,
    "quote_side": slide_quote_side,
    "two_quote": slide_two_quote,
    "table": slide_table,
    "end": slide_end,
}


# ----------------------------------------------------------------- 引用解析
def chapter_key(title: str) -> str:
    m = re.search(r"【(.+?)】", title)
    return m.group(1) if m else title


def clip(text: str, head: int | None, tail: int | None) -> str:
    if head and tail:
        return text[:head].rstrip() + "……" + text[-tail:].lstrip()
    if head:
        return text[:head].rstrip() + "……"
    if tail:
        return "……" + text[-tail:].lstrip()
    return text


def resolve_refs(spec: dict, sections: list[dict], paragraphs: list[dict] | None = None) -> None:
    """把 slides 里的 quote_ref / para / ann 引用换成真实原文。

    引用是刻意保留的间接层：课件规格里不出现长段原文，
    原文一律由本函数从 sections.json / paragraphs.json 取，
    避免手抄出错、也避免模型"顺手改字"。
    """
    by_unit = {}
    by_para = {}
    for sec in sections:
        key = chapter_key(sec["chapter"])
        for unit in sec["units"]:
            by_unit[(key, unit["unit_index"])] = unit
    if paragraphs:
        for p in paragraphs:
            by_para[(chapter_key(p["chapter"]), p["index"])] = p
    else:
        for sec in sections:
            key = chapter_key(sec["chapter"])
            for block in sec["explain_blocks"]:
                by_para[(key, block["index"])] = block

    for slide in spec["slides"]:
        ref = slide.pop("quote_ref", None)
        if ref:
            unit = by_unit[(ref["chapter"], ref["unit"])]
            demo = unit["demo"]
            slide["quote"] = clip(demo["text"], ref.get("head"), ref.get("tail"))
            sources = "、".join(demo["sources"]) if demo["sources"] else "书中引文"
            slide.setdefault("kicker", sources)
            pages = f"p{demo['pages'][0]}" if demo["pages"][0] == demo["pages"][1] else f"p{demo['pages'][0]}–{demo['pages'][1]}"
            slide.setdefault("source", f"《大师们的写作课》{pages}")
        for col in slide.get("columns") or []:
            cref = col.pop("quote_ref", None)
            if cref:
                unit = by_unit[(cref["chapter"], cref["unit"])]
                demo = unit["demo"]
                col["quote"] = clip(demo["text"], cref.get("head"), cref.get("tail"))
                col.setdefault("heading", "、".join(demo["sources"]) or "书中引文")
                pages = (
                    f"p{demo['pages'][0]}"
                    if demo["pages"][0] == demo["pages"][1]
                    else f"p{demo['pages'][0]}–{demo['pages'][1]}"
                )
                # 并置版式同样要给出处页码，学生回头翻原书才找得到
                col.setdefault("source", f"《大师们的写作课》{pages}")
        for block in slide.get("side") or []:
            pref = block.pop("para", None)
            if pref:
                para = by_para[(pref["chapter"], pref["index"])]
                block["text"] = clip(para["text"], pref.get("head"), pref.get("tail"))
            aref = block.pop("ann", None)
            if aref:
                unit = by_unit[(aref["chapter"], aref["unit"])]
                joined = "\n".join(a["text"] for a in unit["annotations"])
                block["text"] = clip(joined, aref.get("head"), aref.get("tail"))
            egroup = block.pop("explains", None)
            if egroup:
                key = egroup["chapter"]
                picked = [by_para[(key, i)] for i in egroup["indexes"] if (key, i) in by_para]
                joined = "\n".join(p["text"] for p in picked)
                block["text"] = clip(joined, egroup.get("head"), egroup.get("tail"))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--sections", default=None, help="sections.json，用于解析 quote_ref / para / ann 引用")
    ap.add_argument("--paragraphs", default=None, help="paragraphs.json，给定后 para 引用可指向任意段落（含批注段）")
    args = ap.parse_args()

    with open(args.spec, encoding="utf-8") as fh:
        spec = json.load(fh)
    if args.sections:
        with open(args.sections, encoding="utf-8") as fh:
            sections = json.load(fh)
        paragraphs = None
        if args.paragraphs:
            with open(args.paragraphs, encoding="utf-8") as fh:
                paragraphs = json.load(fh)
        resolve_refs(spec, sections, paragraphs)

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    page = 0
    for item in spec["slides"]:
        kind = item.get("layout", "bullets")
        if kind not in RENDERERS:
            raise SystemExit(f"未知版式：{kind}")
        page += 1
        RENDERERS[kind](prs, item, page)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    prs.save(args.out)
    counts: dict[str, int] = {}
    for item in spec["slides"]:
        counts[item.get("layout", "bullets")] = counts.get(item.get("layout", "bullets"), 0) + 1
    print(json.dumps({"ok": True, "slides": page, "layouts": counts, "out": os.path.abspath(args.out)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
